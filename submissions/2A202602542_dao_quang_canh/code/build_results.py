"""build_results.py - dựng results.xlsx và các biểu đồ tổng hợp từ log thật.

Chạy từ thư mục bài nộp:  python code/build_results.py
Đọc:  logs/colab_runs/<exp>/seed0/summary.json   (Bước 1-2, Colab)
      logs/kaggle_final/{F01,T00}/seed*/...       (Bước 4, Kaggle)
      eval_out/*.csv                              (kết quả của eval.py trên predictions/)
Số của Bước 3 (suy luận, độ trễ trên Colab) chỉ có trong output của notebook, không có file log:
được chép từ output đó và ghi rõ nguồn ở cột `source`. Độ trễ của Kaggle đọc từ latency_convnext_tiny.json.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import matplotlib
import numpy as np
import pandas as pd

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "code"))
from experiments import BACKBONES, TRAINING  # noqa: E402

COLAB = ROOT / "logs" / "colab_runs"
KAG = ROOT / "logs" / "kaggle_final"
EVAL = ROOT / "eval_out"
CLASS_NAMES = ["Chinee apple", "Lantana", "Parkinsonia", "Parthenium", "Prickly acacia",
               "Rubber vine", "Siam weed", "Snake weed", "Negative"]


def load(p: Path) -> dict:
    return json.loads(p.read_text(encoding="utf8"))


def pm(mean, std, d=4):
    return f"{mean:.{d}f} ± {std:.{d}f}"


# ---------------------------------------------------------------- Backbones
rows = []
for exp, (name, short) in BACKBONES.items():
    s = load(COLAB / exp / "seed0" / "summary.json")
    rows.append({"exp_id": exp, "backbone": s["backbone"], "tag trọng số": s["weight_tag"],
                 "#tham số (M)": round(s["params_m"], 2), "GMAC (224)": round(s["gmacs"], 3), "độ phân giải": 224,
                 "epoch": 12, "seed": 0, "macro-F1 val": s["val_macro_f1"], "top-1 val": s["val_top1"],
                 "train/epoch (s)": round(s["train_time_per_epoch_s"], 1),
                 "độ trễ batch-1 (ms)": None, "epoch tốt nhất": s["best_epoch"],
                 "ghi chú": "1 seed. Độ trễ chỉ đo cho backbone được chọn (xem sheet Latency). "
                            "Tag trọng số khác nhau giữa các backbone nên không phải so sánh thuần kiến trúc."})
bb = pd.DataFrame(rows)

# ---------------------------------------------------------------- Training
t00 = load(COLAB / "T00" / "seed0" / "summary.json")["val_macro_f1"]
rows = []
for exp, (axis, desc, ov) in TRAINING.items():
    s = load(COLAB / exp / "seed0" / "summary.json")
    rows.append({"exp_id": exp, "backbone": "convnext_tiny", "trục": axis, "khác T00 ở điểm nào": desc,
                 "seed": 0, "macro-F1 val": s["val_macro_f1"], "top-1 val": s["val_top1"],
                 "Δ macro-F1 so với T00": s["val_macro_f1"] - t00, "epoch tốt nhất": s["best_epoch"],
                 "ghi chú": ("tổ hợp T04 + T13 (khác T00 hai yếu tố, có chủ đích)" if exp == "T16" else
                             "mốc" if exp == "T00" else "1 seed; |Δ| < ~0,004 không phân biệt được với nhiễu")})
tr = pd.DataFrame(rows)

# nhiễu giữa các seed của mốc (Kaggle, 3 seed): dùng để đọc Δ ở Bước 2
t00_val = [load(KAG / "T00" / f"seed{k}" / "summary.json")["val_macro_f1"] for k in (0, 1, 2)]
noise_std = float(np.std(t00_val, ddof=1))

# ---------------------------------------------------------------- Inference (Bước 3, Colab, T00 seed 0, val)
src3 = "output ô Bước 3 của notebook Colab (checkpoint T00 seed 0, val); không có file log"
lat_c = {  # (config, batch): (p50, p95, p99, img/s) từ output notebook Colab, Tesla T4, 224
    ("fp32", 1): (6.019706, 9.820711, 10.513444, 166.121070), ("fp32", 32): (154.209919, 159.918677, 162.988479, 207.509350),
    ("amp", 1): (10.238367, 13.370161, 14.594689, 97.671831), ("amp", 32): (57.199744, 58.472141, 59.340478, 559.443063),
    ("fp16", 1): (5.620535, 6.322996, 9.381639, 177.918990), ("fp16", 32): (45.167523, 46.409348, 46.719038, 708.473653),
    ("fp32+fusedBN (không áp dụng)", 1): (8.034275, 10.575459, 11.560075, 124.466738),
    ("fp32+fusedBN (không áp dụng)", 32): (147.770958, 152.268489, 154.348355, 216.551347),
    ("TTA lật K=2 (fp32)", 1): (11.847982, 13.546183, 17.449961, 84.402563)}
base_p50 = lat_c[("fp32", 1)][0]
inf = [
    ("I00", "1 view (mốc)", 1, 0.966042, 0.973722, 0.018222, ("fp32", 1)),
    ("I01", "TTA lật ngang (gộp xác suất)", 2, 0.965673, 0.973151, 0.016296, ("TTA lật K=2 (fp32)", 1)),
    ("I03", "TTA lật ngang (gộp logit)", 2, 0.966199, 0.973436, 0.018776, ("TTA lật K=2 (fp32)", 1)),
    ("I02", "5-crop 224 từ 256 (gộp xác suất)", 5, 0.965041, 0.973436, 0.015717, None),
    ("I04a", "độ phân giải test 224", 1, 0.966042, 0.973722, 0.018222, ("fp32", 1)),
    ("I04b", "độ phân giải test 256", 1, 0.966761, 0.974579, 0.017720, None),
    ("I04c", "độ phân giải test 288", 1, 0.975964, 0.981148, 0.014084, None),
    ("I04d", "độ phân giải test 320", 1, 0.973627, 0.978863, 0.015594, None),
    ("I07", "temperature scaling (T=2,107 khớp trên val)", 1, 0.966042, 0.973722, 0.004849, ("fp32", 1)),
]
rows = []
for exp, name, k, f1, t1, ece, lat in inf:
    r = {"exp_id": exp, "phương pháp": name, "mô hình/checkpoint": "T00 seed 0 (Colab)", "K": k,
         "macro-F1 val": f1, "top-1 val": t1, "ECE val": ece}
    if lat:
        p50, p95, p99, ips = lat_c[lat]
        r.update({"p50 batch-1 (ms)": p50, "p95 (ms)": p95, "p99 (ms)": p99, "ảnh/s (batch 1)": ips,
                  "chi phí tương đối so với I00": p50 / base_p50})
    else:
        r.update({"p50 batch-1 (ms)": None, "p95 (ms)": None, "p99 (ms)": None, "ảnh/s (batch 1)": None,
                  "chi phí tương đối so với I00": None})
    r["source"] = src3
    rows.append(r)
inf_df = pd.DataFrame(rows)
inf_df.loc[len(inf_df)] = {"exp_id": "I08", "phương pháp": "gộp BN vào conv", "mô hình/checkpoint": "T00 seed 0 (Colab)",
                           "K": 1, "ghi chú": "Không áp dụng: ConvNeXt dùng LayerNorm, fuse_conv_bn gộp 0 cặp "
                                              "(đã kiểm tra trên ResNet-18/EfficientNet-B0: sai số < 3e-7)",
                           "source": src3}
inf_df.loc[len(inf_df)] = {"exp_id": "I08b", "phương pháp": "fp16 / AMP (xem sheet Latency)", "K": 1,
                           "macro-F1 val": None, "ghi chú": "Chỉ đo độ trễ; độ chính xác fp16/AMP chưa đo riêng",
                           "source": src3}
inf_df.loc[len(inf_df)] = {"exp_id": "I05/I06", "phương pháp": "ensemble, EMA/soup", "ghi chú":
                           "Chưa chạy ở Bước 3. EMA 0,998 được thử như yếu tố huấn luyện (T13) và nằm trong F01."}

# ---------------------------------------------------------------- Final
per_seed = {t: pd.read_csv(EVAL / f"{t}_per_seed.csv") for t in ("F01", "T00")}
rows = []
for t, cfgname in (("F01", "convnext_tiny + TrivialAugment + EMA 0,998; test 288 + temperature scaling"),
                   ("T00", "convnext_tiny + công thức nền; test 224, 1 view (mốc)")):
    for k in (0, 1, 2):
        ps = per_seed[t][per_seed[t]["seed"] == k].iloc[0]
        if t == "F01":
            f1val = load(KAG / "F01" / f"seed{k}" / "final.json")["val_macro_f1"]
        else:
            f1val = load(KAG / "T00" / f"seed{k}" / "summary.json")["val_macro_f1"]
        rows.append({"exp_id": t, "cấu hình": cfgname, "seed": k, "macro-F1 val": f1val,
                     "macro-F1 test": ps["macro_f1"], "top-1 test": ps["top1"], "balanced acc test": ps["balanced_acc"],
                     "ECE test": ps["ece"], "NLL test": ps["nll"], "mean ± std qua seed": ""})
    d = per_seed[t]
    rows.append({"exp_id": t, "cấu hình": "TỔNG HỢP (3 seed, ddof=1)", "seed": "mean ± std",
                 "macro-F1 val": pm(np.mean([r["macro-F1 val"] for r in rows[-3:]]),
                                    np.std([r["macro-F1 val"] for r in rows[-3:]], ddof=1)),
                 "macro-F1 test": pm(d.macro_f1.mean(), d.macro_f1.std(ddof=1)),
                 "top-1 test": pm(d.top1.mean(), d.top1.std(ddof=1)),
                 "balanced acc test": pm(d.balanced_acc.mean(), d.balanced_acc.std(ddof=1)),
                 "ECE test": pm(d.ece.mean(), d.ece.std(ddof=1)), "NLL test": pm(d.nll.mean(), d.nll.std(ddof=1)),
                 "mean ± std qua seed": "3 seed"})
fin = pd.DataFrame(rows)
f_mean, t_mean = per_seed["F01"].macro_f1.mean(), per_seed["T00"].macro_f1.mean()

# ---------------------------------------------------------------- PerClass
pcs = {t: pd.read_csv(EVAL / f"{t}_per_class.csv") for t in ("F01", "T00")}
paper = {"Chinee apple": 0.885, "Snake weed": 0.888}
rows = []
for i, c in enumerate(pcs["F01"]["class"]):
    a, b = pcs["F01"].iloc[i], pcs["T00"].iloc[i]
    rows.append({"lớp": c, "số ảnh test": int(a["support"]),
                 "F01 precision": pm(a.precision_mean, a.precision_std, 3), "F01 recall": pm(a.recall_mean, a.recall_std, 3),
                 "F01 F1": pm(a.f1_mean, a.f1_std, 3),
                 "T00 precision": pm(b.precision_mean, b.precision_std, 3), "T00 recall": pm(b.recall_mean, b.recall_std, 3),
                 "T00 F1": pm(b.f1_mean, b.f1_std, 3),
                 "recall bài báo (ResNet-50, trích dẫn)": paper.get(c)})
pc = pd.DataFrame(rows)

# ---------------------------------------------------------------- Latency
rows = []
for (cfg, bs), (p50, p95, p99, ips) in lat_c.items():
    rows.append({"nguồn": "Colab, Step 3", "cấu hình": cfg, "GPU": "Tesla T4", "dtype": cfg.split()[0].split("+")[0],
                 "độ phân giải": 224, "batch": bs, "gộp BN": "không áp dụng", "p50 (ms)": p50, "p95 (ms)": p95,
                 "p99 (ms)": p99, "ảnh/s": ips})
for r in load(ROOT / "logs" / "kaggle_final" / "latency_convnext_tiny.json"):
    rows.append({"nguồn": "Kaggle, Step 4", "cấu hình": "convnext_tiny " + r["dtype"], "GPU": r["gpu"], "dtype": r["dtype"],
                 "độ phân giải": r["img_size"], "batch": r["batch"], "gộp BN": "không áp dụng",
                 "p50 (ms)": r["p50"], "p95 (ms)": r["p95"], "p99 (ms)": r["p99"], "ảnh/s": r["images_per_s"]})
lat_df = pd.DataFrame(rows)

# ---------------------------------------------------------------- Summary (top 10 theo macro-F1 val)
cand = []
for _, r in bb.iterrows():
    cand.append((r["exp_id"], r["backbone"] + " (công thức nền)", r["macro-F1 val"], r["top-1 val"],
                 f"{r['train/epoch (s)']} s/epoch; {r['GMAC (224)']} GMAC"))
for _, r in tr.iterrows():
    if r["exp_id"] != "T00":
        cand.append((r["exp_id"], "convnext_tiny; " + r["khác T00 ở điểm nào"], r["macro-F1 val"], r["top-1 val"], "1 seed"))
f01v = [load(KAG / "F01" / f"seed{k}" / "final.json") for k in (0, 1, 2)]
cand.append(("F01 (val, 288 + TS)", "T16 + test ở 288 + temperature; mean 3 seed",
             float(np.mean([v["val_macro_f1"] for v in f01v])), float(np.mean([v["val_top1"] for v in f01v])),
             "độ trễ p95 batch-1 fp32 ở 288: xem Latency"))
cand.append(("I04c (val, 288)", "T00 seed 0 test ở 288", 0.975964, 0.981148, "p50 8,3 ms @288 (Kaggle)"))
cand.sort(key=lambda x: -x[2])
summ = pd.DataFrame([{"hạng": i + 1, "exp_id": c[0], "cấu hình": c[1], "macro-F1 val": c[2], "top-1 val": c[3],
                      "chi phí / độ trễ": c[4]} for i, c in enumerate(cand[:10])])
head = pd.DataFrame([
    {"hạng": "", "exp_id": "CHUNG KẾT (test, 3 seed)", "cấu hình": "F01 so với mốc T00",
     "macro-F1 val": "", "top-1 val": "",
     "chi phí / độ trễ": f"F01 macro-F1 test {pm(per_seed['F01'].macro_f1.mean(), per_seed['F01'].macro_f1.std(ddof=1))} "
                         f"so với T00 {pm(t_mean, per_seed['T00'].macro_f1.std(ddof=1))}; Δ = {f_mean - t_mean:+.4f}"}])
summ = pd.concat([head, summ], ignore_index=True)
notes = pd.DataFrame({"ghi chú": [
    f"Nhiễu giữa các seed của mốc T00 (Kaggle, val macro-F1, 3 seed): std = {noise_std:.4f}.",
    "Các Δ của Bước 2 (1 seed) nhỏ hơn khoảng 0,004 không phân biệt được với nhiễu.",
    "F01 và mốc T00 ở Bước 4 đều chạy trên Kaggle (T4). Mốc T00 của Bước 2 chạy trên Colab (macro-F1 val 0,9660).",
    "Lần chạy đầu của Bước 4 dựng nhầm mốc T00 với backbone resnet50 (lỗi, đã sửa); dự đoán đó bị loại, "
    "xem predictions_discarded_resnet50_baseline/.",
    "Số tham chiếu bài báo (95,7%, 88,5%, 88,8%) là số trích dẫn, không phải kết quả của bài này."]})

# ---------------------------------------------------------------- ghi xlsx
out = ROOT / "results.xlsx"
with pd.ExcelWriter(out, engine="openpyxl") as xw:
    for name, df in (("Summary", summ), ("Backbones", bb), ("Training", tr), ("Inference", inf_df), ("Final", fin),
                     ("PerClass", pc), ("Latency", lat_df)):
        df.to_excel(xw, sheet_name=name, index=False)
    notes.to_excel(xw, sheet_name="Summary", index=False, startrow=len(summ) + 3)
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    for ws in xw.book.worksheets:
        ws.freeze_panes = "A2"
        for c in ws[1]:
            c.font = Font(bold=True)
        for i, col in enumerate(ws.columns, 1):
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col[:40])
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, float):
                    c.number_format = "0.0000"
    # tô nổi bật dòng tốt nhất
    hl = PatternFill("solid", fgColor="FFF2CC")
    for sheet, col in (("Backbones", "macro-F1 val"), ("Training", "macro-F1 val")):
        ws = xw.book[sheet]
        idx = [c.value for c in ws[1]].index(col) + 1
        best = max(range(2, ws.max_row + 1), key=lambda r: ws.cell(r, idx).value)
        for c in ws[best]:
            c.fill = hl
print("đã ghi", out)

# ---------------------------------------------------------------- biểu đồ
fig_dir = ROOT / "figures"
fig_dir.mkdir(exist_ok=True)

cm = pd.read_csv(EVAL / "F01_confusion_sum.csv", index_col=0).to_numpy()
fig, ax = plt.subplots(figsize=(7.5, 6.5))
norm = cm / cm.sum(1, keepdims=True)
im = ax.imshow(norm, cmap="Blues", vmin=0, vmax=1)
for i in range(9):
    for j in range(9):
        ax.text(j, i, cm[i, j], ha="center", va="center", fontsize=7, color="white" if norm[i, j] > 0.5 else "black")
ax.set_xticks(range(9), CLASS_NAMES, rotation=45, ha="right")
ax.set_yticks(range(9), CLASS_NAMES)
ax.set(xlabel="dự đoán", ylabel="nhãn thật", title="F01: ma trận nhầm lẫn trên test (cộng 3 seed; màu = tỉ lệ theo hàng)")
fig.colorbar(im, fraction=0.04)
fig.tight_layout()
fig.savefig(fig_dir / "F01_confusion_matrix.png", dpi=140)
plt.close(fig)

# top nhầm lẫn
off = [(cm[i, j], CLASS_NAMES[i], CLASS_NAMES[j], cm[i, j] / cm[i].sum()) for i in range(9) for j in range(9) if i != j]
off.sort(reverse=True)
pd.DataFrame(off[:10], columns=["số ảnh (3 seed)", "nhãn thật", "dự đoán", "tỉ lệ trong lớp thật"]).to_csv(
    EVAL / "F01_top_confusions.csv", index=False)

# đánh đổi độ chính xác - độ trễ (Bước 3)
pts = [("I00 1-view", 0.966042, 6.02), ("I01 TTA lật K=2", 0.965673, 11.85), ("fp16 1-view", 0.966042, 5.62),
       ("I04c test 288 (Kaggle p50)", 0.975964, 8.3)]
fig, ax = plt.subplots(figsize=(6.5, 4.2))
for n, f, l in pts:
    ax.scatter(l, f, s=50)
    ax.annotate(n, (l, f), textcoords="offset points", xytext=(5, 4), fontsize=8)
ax.set(xlabel="độ trễ p50, batch 1 (ms, Tesla T4)", ylabel="macro-F1 val", title="Đánh đổi độ chính xác - độ trễ (ConvNeXt-tiny, T00)")
ax.grid(alpha=0.3)
fig.tight_layout()
fig.savefig(fig_dir / "inference_tradeoff.png", dpi=140)
plt.close(fig)

# macro-F1 val theo backbone
fig, ax = plt.subplots(figsize=(7, 3.8))
order = bb.sort_values("macro-F1 val")
ax.barh(order["backbone"], order["macro-F1 val"])
for y, v in enumerate(order["macro-F1 val"]):
    ax.text(v + 0.003, y, f"{v:.3f}", va="center", fontsize=8)
ax.set(xlim=(0.7, 1.0), xlabel="macro-F1 val (1 seed, công thức nền)", title="Bước 1: so sánh backbone")
fig.tight_layout()
fig.savefig(fig_dir / "B_backbones_val_f1.png", dpi=140)
plt.close(fig)

# Δ Bước 2 so với T00 (bỏ T01, T02 vì quá lớn)
d = tr[~tr["exp_id"].isin(["T00", "T01", "T02"])]
fig, ax = plt.subplots(figsize=(7.5, 4.2))
ax.bar(d["exp_id"], d["Δ macro-F1 so với T00"])
ax.axhline(0, color="k", lw=0.8)
ax.axhspan(-noise_std * 2, noise_std * 2, color="gray", alpha=0.2, label=f"±2 std của mốc ({noise_std:.4f})")
ax.set(ylabel="Δ macro-F1 val so với T00", title="Bước 2: Δ so với T00 (1 seed)")
ax.legend(fontsize=8)
fig.tight_layout()
fig.savefig(fig_dir / "T_delta_vs_T00.png", dpi=140)
plt.close(fig)
print("đã ghi biểu đồ vào", fig_dir)
